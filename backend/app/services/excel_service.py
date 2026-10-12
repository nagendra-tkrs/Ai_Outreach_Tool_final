import io
import os
import csv
import re
import openpyxl
from app.utils.validators import (
    REQUIRED_COLUMNS,
    resolve_header,
    normalize_website,
    normalize_company_name,
    validate_row,
    suggested_fixes,
    resolve_service,
)
from app.config.database import get_db

FIELD_MAP = {
    'Company Name': 'companyName',
    'Website': 'website',
    'Evidence': 'evidence',
    'Recommended Service': 'recommendedService',
    'Industry': 'industry',
    'Country': 'country',
    'Contact Name': 'contactName',
    'Contact Email': 'contactEmail',
    'Revenue': 'revenue',
    'Employee Count': 'employeeCount',
    'Notes': 'notes',
    'Source': 'source'
}

SUPPORTED_EXTENSIONS = {'.xlsx', '.xls', '.csv'}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB, mirrors the frontend text

TEMPLATE_HEADERS = [
    'Company Name', 'Website', 'Evidence', 'Recommended Service',
    'Industry', 'Country', 'Contact Name', 'Contact Email',
    'Revenue', 'Employee Count', 'Notes', 'Source'
]

TEMPLATE_SAMPLE_ROWS = [
    [
        "Acme Corp", "https://acme.com",
        "Acme is scaling logistics with AI-ready data pipelines and automation gaps.",
        "AI Strategy Consulting", "Logistics", "USA",
        "John Doe; Sarah Lee", "john@acme.com; sarah@acme.com",
        "$10M-$50M", "200-500", "Referral from regional conference", "LinkedIn"
    ],
    [
        "Nova Health", "https://novah.com",
        "Nova is expanding predictive analytics but has no structured AI roadmap.",
        "AI Readiness and Opportunity Assessments", "Healthcare", "UK",
        "Jane Smith", "jane@novah.co.uk",
        "$5M-$10M", "50-200", "", "Trade show"
    ],
    [
        "Bright Manufacturing", "https://brightmfg.com",
        "Bright wants to modernize production lines with AI and automation workshops.",
        "AI & Automation Workshops", "Manufacturing", "Germany",
        "", "",
        "$50M+", "1000+", "Requested at CTO roundtable", "Referral"
    ],
]


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_xlsx(content: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    if not wb.sheetnames:
        raise ValueError('Excel file contains no sheets')
    sheet = wb[wb.sheetnames[0]]
    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        raise ValueError('Excel file is empty or has no data rows')
    headers = [_cell_to_str(h) for h in header_row]
    data = []
    for row in rows_iter:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        data.append([_cell_to_str(cell) for cell in row])
    return headers, data


def _parse_xls(content: bytes):
    try:
        import xlrd
    except ImportError:  # pragma: no cover
        raise ValueError('XLS support requires the "xlrd" package. Install with: pip install xlrd')
    book = xlrd.open_workbook(file_contents=content)
    if book.nsheets == 0:
        raise ValueError('Excel file contains no sheets')
    sheet = book.sheet_by_index(0)
    if sheet.nrows == 0:
        raise ValueError('Excel file is empty or has no data rows')
    headers = [_cell_to_str(sheet.cell_value(0, c)) for c in range(sheet.ncols)]
    data = []
    for r in range(1, sheet.nrows):
        row = [_cell_to_str(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
        if all(cell == "" for cell in row):
            continue
        data.append(row)
    return headers, data


def _parse_csv(content: bytes):
    text = content.decode('utf-8-sig', errors='replace')
    reader = csv.reader(io.StringIO(text))
    try:
        header_row = next(reader, None)
    except StopIteration:
        header_row = None
    if not header_row:
        raise ValueError('CSV file is empty or has no data rows')
    headers = [_cell_to_str(h) for h in header_row]
    data = []
    for row in reader:
        if all(_cell_to_str(cell) == "" for cell in row):
            continue
        data.append([_cell_to_str(cell) for cell in row])
    return headers, data


def parse_excel_file(file_bytes: bytes, filename: str):
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        display = ext if ext else "(no extension)"
        raise ValueError(
            f"Unsupported file format '{display}'. Please upload a .xlsx, .xls, or .csv file."
        )
    if ext == '.csv':
        headers, data = _parse_csv(file_bytes)
        sheet_name = 'CSV'
    elif ext == '.xls':
        headers, data = _parse_xls(file_bytes)
        sheet_name = 'XLS'
    else:
        headers, data = _parse_xlsx(file_bytes)
        sheet_name = 'Sheet1'

    return {"data": data, "sheetName": sheet_name, "headers": headers}


def validate_headers(headers: list[str]):
    """Ensure every required column is present after flexible header mapping."""
    resolved = {resolve_header(h) for h in headers}
    resolved.discard(None)
    missing = [col for col in REQUIRED_COLUMNS if col not in resolved]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")


def normalize_row(raw_row: dict) -> dict:
    """Map a dict keyed by canonical headers into the normalized internal schema."""
    normalized = {}
    for header, field in FIELD_MAP.items():
        normalized[field] = raw_row.get(header, "")
    normalized['companyName'] = normalize_company_name(normalized.get('companyName', ''))
    normalized['website'] = (normalized.get('website') or '').strip()
    normalized['normalizedWebsite'] = normalize_website(normalized['website'])
    normalized['evidence'] = (normalized.get('evidence') or '').strip()
    normalized['recommendedService'] = (normalized.get('recommendedService') or '').strip()
    return normalized


def split_contacts(value: str) -> list[str]:
    """Split a multi-contact cell on ';' or '|', trimming and dropping empties."""
    if not value:
        return []
    parts = re.split(r'[;|]', value)
    return [p.strip() for p in parts if p and p.strip()]


async def load_existing_websites(db) -> set:
    """Set of normalized bare domains already in the companies DB.

    Reads all three domain-ish fields so legacy rows with 'https://...'
    normalizedWebsite values still deduplicate correctly.
    """
    existing = set()
    if db is None:
        return existing
    cursor = db.companies.find({}, {"normalizedWebsite": 1, "website": 1, "officialDomain": 1})
    async for comp in cursor:
        for field in ("normalizedWebsite", "website", "officialDomain"):
            value = comp.get(field)
            if value:
                domain = normalize_website(value)
                if domain:
                    existing.add(domain)
    return existing


async def process_excel(file_bytes: bytes, filename: str) -> dict:
    parsed = parse_excel_file(file_bytes, filename)
    raw_headers = parsed["headers"]
    raw_data = parsed["data"]

    validate_headers(raw_headers)

    # Resolve each column once: index -> canonical header (unknown columns dropped)
    col_map = {}
    for idx, header in enumerate(raw_headers):
        canonical = resolve_header(header)
        if canonical:
            col_map[idx] = canonical

    data = []
    for row in raw_data:
        row_dict = {}
        for idx, canonical in col_map.items():
            if idx < len(row):
                row_dict[canonical] = row[idx]
        data.append(row_dict)

    if not data:
        raise ValueError('Excel file is empty or has no data rows')

    normalized_rows = []
    for idx, row in enumerate(data):
        norm = normalize_row(row)
        norm['rowNumber'] = idx + 2

        # --- Recommended Service mapping (log original vs mapped) ---
        raw_service = norm.get('recommendedService', '')
        canonical_service, was_mapped = resolve_service(raw_service)
        norm['originalService'] = raw_service
        norm['mappedService'] = canonical_service if was_mapped else ""
        norm['serviceMapped'] = bool(was_mapped)
        if canonical_service:
            norm['recommendedService'] = canonical_service

        # --- Contacts: store raw (never reject) + parsed lists ---
        norm['contactNames'] = split_contacts(norm.get('contactName', ""))
        norm['contactEmails'] = split_contacts(norm.get('contactEmail', ""))

        normalized_rows.append(norm)

    existing_website_set = await load_existing_websites(get_db())

    processed_rows = []
    seen_names = {}
    seen_websites = {}

    for row in normalized_rows:
        errors = validate_row(row)
        status = 'invalid' if errors else 'valid'

        if status == 'valid':
            name_lower = row['companyName'].lower()
            if name_lower in seen_names:
                status = 'duplicate_in_file'
                errors.append(f"Duplicate company name (also row {seen_names[name_lower]})")
            else:
                seen_names[name_lower] = row['rowNumber']

        if status == 'valid':
            norm_web = row['normalizedWebsite']
            if norm_web in seen_websites:
                status = 'duplicate_in_file'
                errors.append(f"Duplicate website (also row {seen_websites[norm_web]})")
            else:
                seen_websites[norm_web] = row['rowNumber']

        if status == 'valid' and row['normalizedWebsite'] in existing_website_set:
            status = 'existing_in_db'
            errors.append('Company already exists in database')

        row_data = dict(row)
        row_data['status'] = status
        row_data['validationErrors'] = errors
        row_data['suggestedFix'] = suggested_fixes(errors, row)
        processed_rows.append(row_data)

    summary = {
        'total': len(processed_rows),
        'valid': len([r for r in processed_rows if r['status'] == 'valid']),
        'invalid': len([r for r in processed_rows if r['status'] == 'invalid']),
        'duplicates': len([r for r in processed_rows if r['status'] == 'duplicate_in_file']),
        'existingInDb': len([r for r in processed_rows if r['status'] == 'existing_in_db'])
    }

    return {
        'originalFilename': filename,
        'summary': summary,
        'rows': processed_rows
    }


# ---- Report / template builders ---------------------------------------------


def build_import_template() -> bytes:
    """Return TKSR_Import_Template.xlsx bytes (12 headers + 3 sample rows)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Import Template"
    ws.append(TEMPLATE_HEADERS)
    for row in TEMPLATE_SAMPLE_ROWS:
        ws.append(row)
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    widths = {"A": 22, "B": 24, "C": 70, "D": 36, "E": 16, "F": 12,
              "G": 22, "H": 22, "I": 12, "J": 12, "K": 34, "L": 12}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_invalid_rows_report(rows: list) -> bytes:
    """Return invalid_rows.xlsx bytes: Row Number | Company Name | Status | Error | Suggested Fix."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invalid Rows"
    headers = ["Row Number", "Company Name", "Status", "Error", "Suggested Fix"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)

    for row in rows:
        status = row.get('status', '')
        if status == 'valid':
            continue
        errors = row.get('validationErrors') or []
        fixes = row.get('suggestedFix') or []
        ws.append([
            row.get('rowNumber', ''),
            row.get('companyName', ''),
            status,
            "\n".join(errors) if errors else "",
            "\n".join(fixes) if fixes else "",
        ])

    for col, width in {"A": 12, "B": 26, "C": 20, "D": 60, "E": 60}.items():
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()