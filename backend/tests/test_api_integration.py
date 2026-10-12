import io
import pytest
import openpyxl
from bson import ObjectId
from datetime import datetime, timezone
import httpx

from app.main import app
from app.config.database import connect_db, close_db, get_db

@pytest.fixture(scope="session", autouse=True)
async def setup_database():
    await connect_db()
    yield
    await close_db()

@pytest.fixture
async def client():
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as c:
        yield c

async def test_health_check(client):
    res = await client.get("/api/health")
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert "running" in data["message"].lower()

async def test_settings_endpoints(client):
    # GET settings
    res = await client.get("/api/settings")
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert "ai" in data["data"]
    assert "smtp" in data["data"]

    # POST update settings
    payload = {
        "ai": {"provider": "ollama", "model": "llama3"},
        "smtp": {"fromName": "Test Outreach"},
        "senderProfile": {"senderName": "Agent Test"}
    }
    update_res = await client.post("/api/settings", json=payload)
    assert update_res.status_code == 200
    assert update_res.json()["success"] is True

    # Check updated settings
    verify_res = await client.get("/api/settings")
    vdata = verify_res.json()["data"]
    assert vdata["ai"]["provider"] == "ollama"
    assert vdata["smtp"]["fromName"] == "Test Outreach"

async def test_company_and_contact_pattern_flow(client):
    db = get_db()
    now = datetime.now(timezone.utc)

    # Clean any leftover test records
    await db.companies.delete_many({"normalizedWebsite": "patterntech.com"})
    await db.contacts.delete_many({"email": {"$regex": "@patterntech.com"}})

    # 1. Create company with known contacts
    test_comp = {
        "companyName": "PatternTech Enterprises",
        "website": "https://patterntech.com",
        "normalizedWebsite": "patterntech.com",
        "officialDomain": "patterntech.com",
        "corporateEmailDomain": "patterntech.com",
        "createdAt": now,
        "updatedAt": now,
    }
    insert_res = await db.companies.insert_one(test_comp)
    comp_id = str(insert_res.inserted_id)

    # Insert verified contacts
    await db.contacts.insert_many([
        {
            "companyId": insert_res.inserted_id,
            "name": "Sarah Connor",
            "title": "VP Engineering",
            "roleCategory": "vp",
            "email": "sarah.connor@patterntech.com",
            "personVerified": "true",
            "companyMatchStatus": "VERIFIED",
            "emailStatus": "DIRECTLY_VERIFIED",
            "createdAt": now,
            "updatedAt": now,
        },
        {
            "companyId": insert_res.inserted_id,
            "name": "John Connor",
            "title": "Head of Data",
            "roleCategory": "head_of_data",
            "email": "john.connor@patterntech.com",
            "personVerified": "true",
            "companyMatchStatus": "VERIFIED",
            "emailStatus": "DIRECTLY_VERIFIED",
            "createdAt": now,
            "updatedAt": now,
        }
    ])

    try:
        # Test GET pattern
        pattern_res = await client.get(f"/api/companies/{comp_id}/contacts/pattern")
        assert pattern_res.status_code == 200
        pdata = pattern_res.json()
        assert pdata["success"] is True
        assert pdata["data"]["pattern"] == "first.last"
        assert pdata["data"]["emailDomain"] == "patterntech.com"
        assert pdata["data"]["confidence"] >= 80

        # Test apply pattern to a new contact without email
        contact_unverified = await db.contacts.insert_one({
            "companyId": insert_res.inserted_id,
            "name": "Kyle Reese",
            "title": "Operations Director",
            "roleCategory": "director",
            "email": "",
            "createdAt": now,
            "updatedAt": now,
        })
        kyle_id = str(contact_unverified.inserted_id)

        apply_res = await client.post(
            f"/api/companies/{comp_id}/contacts/apply-pattern",
            json={"pattern": "first.last"}
        )
        assert apply_res.status_code == 200
        assert apply_res.json()["success"] is True

        # Check that Kyle Reese got email generated
        updated_kyle = await db.contacts.find_one({"_id": ObjectId(kyle_id)})
        assert updated_kyle["email"] == "kyle.reese@patterntech.com"
        assert updated_kyle["emailStatus"] == "PATTERN_BASED_UNVERIFIED"

    finally:
        # Clean up
        await db.companies.delete_one({"_id": insert_res.inserted_id})
        await db.contacts.delete_many({"companyId": insert_res.inserted_id})

async def test_research_endpoints(client):
    db = get_db()
    now = datetime.now(timezone.utc)

    # Test invalid company ID
    res_invalid = await client.get("/api/companies/invalid-id-12345/research")
    assert res_invalid.status_code == 400
    assert "Invalid company ID" in res_invalid.json()["message"]

    # Test non-existent company ID
    fake_id = str(ObjectId())
    res_404 = await client.get(f"/api/companies/{fake_id}/research")
    assert res_404.status_code == 404
    assert "Company not found" in res_404.json()["message"]

    # Clean any leftover records
    await db.companies.delete_many({"normalizedWebsite": "researchtest.com"})

    # Create real company
    c_res = await db.companies.insert_one({
        "companyName": "Research Test Inc",
        "website": "https://researchtest.com",
        "normalizedWebsite": "researchtest.com",
        "createdAt": now,
        "updatedAt": now,
    })
    comp_id = str(c_res.inserted_id)

    try:
        # Before research run: should return 200 with data: null
        res_empty = await client.get(f"/api/companies/{comp_id}/research")
        assert res_empty.status_code == 200
        assert res_empty.json()["data"] is None

        # Insert research record
        await db.companyresearches.insert_one({
            "companyId": c_res.inserted_id,
            "summary": "Research Test Inc summary information",
            "valueProposition": "Value prop",
            "coreOfferings": ["AI Consulting", "ML Ops"],
            "targetAudience": "Enterprise",
            "icpFitScore": 90,
            "fallbackUsed": False,
            "createdAt": now,
            "updatedAt": now,
        })

        # After research record exists: should return 200 with payload
        res_filled = await client.get(f"/api/companies/{comp_id}/research")
        assert res_filled.status_code == 200
        assert res_filled.json()["data"]["summary"] == "Research Test Inc summary information"
        assert res_filled.json()["data"]["icpFitScore"] == 90

    finally:
        # Clean up
        await db.companies.delete_one({"_id": c_res.inserted_id})
        await db.companyresearches.delete_many({"companyId": c_res.inserted_id})

async def test_excel_import_and_company_listing(client):
    db = get_db()
    await db.companies.delete_many({"companyName": {"$in": ["Acme Robotics", "Nova Health"]}})

    # Create in-memory excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Company Name", "Website", "Evidence", "Recommended Service", "Industry", "Country"])
    ws.append(["Acme Robotics", "https://acmerobotics.com", "Expanding robotic supply chain operations.", "AI & Automation Workshops", "Robotics", "USA"])
    ws.append(["Nova Health", "https://novahealth.co.uk", "Deploying healthcare intelligence models.", "AI Strategy Consulting", "Healthcare", "UK"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    files = {
        "file": ("test_import.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }

    res_upload = await client.post("/api/import/excel", files=files)
    assert res_upload.status_code == 200
    up_data = res_upload.json()
    assert up_data["success"] is True
    batch_id = up_data["data"]["batchId"]
    assert up_data["data"]["summary"]["total"] == 2

    try:
        # Confirm import
        res_confirm = await client.post(f"/api/import/{batch_id}/confirm")
        assert res_confirm.status_code == 200
        conf_data = res_confirm.json()
        assert conf_data["success"] is True

        # List companies
        res_list = await client.get("/api/companies", params={"search": "Acme Robotics"})
        assert res_list.status_code == 200
        list_data = res_list.json()
        assert list_data["success"] is True
        assert any(c["companyName"] == "Acme Robotics" for c in list_data["data"]["companies"])

    finally:
        # Clean up created records
        await db.importbatches.delete_one({"_id": ObjectId(batch_id)})
        await db.companies.delete_many({"companyName": {"$in": ["Acme Robotics", "Nova Health"]}})

async def test_outreach_endpoints(client):
    db = get_db()
    now = datetime.now(timezone.utc)

    # Clean previous records
    await db.companies.delete_many({"normalizedWebsite": "outreachtech.com"})

    comp_doc = await db.companies.insert_one({
        "companyName": "OutreachTech",
        "website": "https://outreachtech.com",
        "normalizedWebsite": "outreachtech.com",
        "evidence": "OutreachTech is deploying automated sales pipelines.",
        "recommendedService": "AI Strategy Consulting",
        "createdAt": now,
        "updatedAt": now,
    })
    comp_id = str(comp_doc.inserted_id)

    try:
        # Generate outreach sequence
        gen_res = await client.post(
            f"/api/companies/{comp_id}/outreach/generate",
            json={"service": "AI Consulting", "tone": "Professional"}
        )
        assert gen_res.status_code == 200
        gen_data = gen_res.json()
        assert gen_data["success"] is True
        assert len(gen_data["data"]) == 3  # 3-step sequence

        # Fetch outreach
        list_res = await client.get(f"/api/companies/{comp_id}/outreach")
        assert list_res.status_code == 200
        assert len(list_res.json()["data"]) == 3

        # Check global stats
        stats_res = await client.get("/api/outreach/stats")
        assert stats_res.status_code == 200
        assert stats_res.json()["success"] is True

        # Check global records
        records_res = await client.get("/api/outreach/records")
        assert records_res.status_code == 200
        assert records_res.json()["success"] is True

        # Check Excel export endpoint returns valid XLSX binary bytes
        export_res = await client.get("/api/outreach/export")
        assert export_res.status_code == 200
        assert export_res.headers.get("content-type") == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert 'filename="outreach_report_' in export_res.headers.get("content-disposition", "")
        assert export_res.content.startswith(b"PK\x03\x04")
        exported_wb = openpyxl.load_workbook(io.BytesIO(export_res.content))
        assert "Outreach Tracking" in exported_wb.sheetnames
        assert exported_wb["Outreach Tracking"].max_column == 15

    finally:
        await db.companies.delete_one({"_id": comp_doc.inserted_id})
        await db.outreachmessages.delete_many({"companyId": comp_doc.inserted_id})
        await db.outreachrecords.delete_many({"companyId": comp_doc.inserted_id})


